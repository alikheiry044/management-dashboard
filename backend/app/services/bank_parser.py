"""
پارسر فایل «مرور بانک‌ها (دریافت و پرداخت)».

ستون‌های مورد انتظار (ردیف عنوان می‌تواند هر جای ۳۰ ردیف اول باشد):
    کد تفصیلی | عنوان | اول دوره | دریافت‌ها | پرداخت‌ها | مانده

- هم .xls قدیمی (xlrd) و هم .xlsx (openpyxl) پشتیبانی می‌شود.
- حروف عربی (ي ك) به فارسی (ی ک) تبدیل می‌شوند.
- «-» در ستون مبلغ یعنی صفر.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

_PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


class BankParseError(ValueError):
    """خطای قابل نمایش به کاربر (پیام فارسی)."""


@dataclass
class BankRow:
    detail_code: str
    title: str
    account_number: str | None
    opening_balance: int
    receipts: int
    payments: int
    balance: int


# ---------- helpers ----------
def fix_text(value: Any) -> str:
    """نرمال‌سازی متن: ي→ی، ك→ک، نیم‌فاصله→فاصله، فاصله‌های اضافه حذف."""
    if value is None:
        return ""
    text = str(value).replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ")
    return re.sub(r"\s+", " ", text).strip()


def _key(value: Any) -> str:
    """کلید مقایسه‌ی عنوان ستون: بدون فاصله."""
    return fix_text(value).replace(" ", "")


def to_int(value: Any) -> int:
    """عدد را از سلول بخوان؛ خالی و «-» یعنی ۰؛ پرانتز یا منفی انتهایی یعنی منفی."""
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(round(value))
    text = fix_text(value).translate(_PERSIAN_DIGITS)
    text = text.replace(",", "").replace("٬", "").replace(" ", "")
    if text in {"", "-", "--"}:
        return 0
    negative = False
    if text.startswith("(") and text.endswith(")"):
        negative, text = True, text[1:-1]
    if text.endswith("-"):
        negative, text = True, text[:-1]
    try:
        number = int(round(float(text)))
    except ValueError:
        raise BankParseError(f"مقدار «{text}» عدد معتبری نیست")
    return -number if negative else number


def clean_code(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return fix_text(value).translate(_PERSIAN_DIGITS)


_ACCOUNT_RE = re.compile(r"(?:^|\s)(\d[\d.\-/]{5,}\d?)\s*$")


def split_account(title: str) -> tuple[str, str | None]:
    """اگر آخر عنوان یک شماره حساب بلند هست، جدایش کن."""
    match = _ACCOUNT_RE.search(title.translate(_PERSIAN_DIGITS))
    if not match:
        return title, None
    account = match.group(1)
    name = title[: len(title) - len(match.group(0))].strip(" -")
    return (name or title), account


# ---------- reading the file ----------
def _read_rows(content: bytes) -> list[list[Any]]:
    head = content[:8]
    if head.startswith(b"\xd0\xcf\x11\xe0"):  # .xls (OLE2)
        try:
            import xlrd
        except ImportError:
            raise BankParseError("برای خواندن فایل .xls باید xlrd نصب شود: pip install xlrd")
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_index(0)
        return [sheet.row_values(i) for i in range(sheet.nrows)]

    if head.startswith(b"PK"):  # .xlsx
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]

    # بعضی نرم‌افزارهای حسابداری فایل HTML با پسوند xls می‌سازند
    try:
        import pandas as pd
        tables = pd.read_html(io.BytesIO(content))
        if tables:
            return tables[0].fillna("").values.tolist()
    except Exception:
        pass
    raise BankParseError("فرمت فایل شناخته نشد. فایل باید اکسل (.xls یا .xlsx) باشد.")


_COLUMNS = {
    "code": {"کدتفصیلی", "کدتفضیلی"},
    "title": {"عنوان", "شرح", "نام"},
    "opening": {"اولدوره", "ماندهاولدوره"},
    "receipts": {"دریافتها", "دریافت"},
    "payments": {"پرداختها", "پرداخت"},
    "balance": {"مانده"},
}


def _find_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    for idx, row in enumerate(rows[:30]):
        keys = [_key(c) for c in row]
        found: dict[str, int] = {}
        for name, aliases in _COLUMNS.items():
            for col, k in enumerate(keys):
                if k in aliases and name not in found:
                    found[name] = col
        if {"code", "title", "receipts", "payments", "balance"} <= found.keys():
            return idx, found
    raise BankParseError(
        "ستون‌های «کد تفصیلی، عنوان، دریافت‌ها، پرداخت‌ها، مانده» پیدا نشد. "
        "مطمئن شوید فایل «مرور بانک‌ها (دریافت و پرداخت)» را انتخاب کرده‌اید."
    )


# ---------- public API ----------
def parse_bank_accounts(content: bytes) -> list[BankRow]:
    """از بایت‌های فایل (xls/xlsx) مستقیم بخوان."""
    return parse_bank_rows(_read_rows(content))


def parse_bank_rows(rows: list[list[Any]]) -> list[BankRow]:
    """از ردیف‌های آماده بخوان (ردیف عنوان می‌تواند هر جای ۳۰ ردیف اول باشد)."""
    header_idx, cols = _find_header(rows)

    def cell(row: list[Any], name: str) -> Any:
        col = cols.get(name)
        return row[col] if col is not None and col < len(row) else None

    result: list[BankRow] = []
    seen: set[str] = set()
    for row in rows[header_idx + 1:]:
        code = clean_code(cell(row, "code"))
        title = fix_text(cell(row, "title"))
        if not code:
            continue  # ردیف خالی یا ردیف «جمع کل»
        if code in seen:
            raise BankParseError(f"کد تفصیلی «{code}» در فایل تکراری است")
        seen.add(code)

        _, account = split_account(title)
        result.append(BankRow(
            detail_code=code,
            title=title,
            account_number=account,
            opening_balance=to_int(cell(row, "opening")),
            receipts=to_int(cell(row, "receipts")),
            payments=to_int(cell(row, "payments")),
            balance=to_int(cell(row, "balance")),
        ))

    if not result:
        raise BankParseError("هیچ ردیف داده‌ای در فایل پیدا نشد")
    return result
